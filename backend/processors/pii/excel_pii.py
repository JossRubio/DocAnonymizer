"""Anonimizacion de libros .xlsx conservando el texto no sensible.

Este procesador es deliberadamente mas conservador que los otros dos, porque en
una hoja de calculo un cambio de tipo rompe cosas que no se ven:

  - Las formulas no se tocan nunca. Sustituir una deja el libro con #REF! y se
    lleva por delante todo lo que dependiera de ella.
  - Las celdas numericas dependen del nivel elegido. Escribir texto donde habia
    un numero vacia los graficos que lo referencian y rompe los calculos, asi
    que en nivel suave no se tocan y en nivel intermedio solo si la cabecera de
    su columna indica que la columna contiene datos sensibles.
  - Una hoja solo se renombra si su nombre contiene datos personales y ninguna
    formula la referencia: openpyxl no reescribe las referencias al renombrar.

La cabecera de columna se propaga como contexto a las celdas de esa columna. Es
lo que permite anonimizar una columna "Salario" o "RUT" cuyo encabezado esta
decenas de filas mas arriba, fuera del alcance de cualquier ventana de contexto.
"""

import re

import openpyxl

from .engine import PiiEngine
from .levels import (NUMERIC_ALWAYS, NUMERIC_BY_HEADER, NUMERIC_POLICY,
                     level_value)
from .patterns import SENSITIVE_HEADER
from .runs import apply_to_text

# Cuantas filas se inspeccionan buscando la fila de cabecera. En los libros
# reales la cabecera no siempre esta en la fila 1: suele haber un titulo encima.
_HEADER_SEARCH_ROWS = 8
_MAX_SHEET_TITLE = 31


def _is_formula(cell) -> bool:
    if getattr(cell, "data_type", None) == "f":
        return True
    value = cell.value
    return isinstance(value, str) and value.startswith("=")


def _find_header_row(ws) -> int:
    """Primera fila que parece una cabecera: varias celdas de texto y ningun numero.

    Devuelve 0 si no encuentra ninguna, en cuyo caso no se propaga contexto.
    """
    max_row = min(ws.max_row or 0, _HEADER_SEARCH_ROWS)
    for row_idx in range(1, max_row + 1):
        texts = 0
        numbers = 0
        for cell in ws[row_idx]:
            value = cell.value
            if value is None:
                continue
            if isinstance(value, bool):
                continue
            if isinstance(value, (int, float)):
                numbers += 1
            elif str(value).strip():
                texts += 1
        if texts >= 3 and numbers == 0:
            return row_idx
    return 0


def _column_headers(ws) -> dict:
    """Texto de la cabecera de cada columna, indexado por numero de columna."""
    header_row = _find_header_row(ws)
    if not header_row:
        return {}
    headers = {}
    for cell in ws[header_row]:
        if cell.value is not None and str(cell.value).strip():
            headers[cell.column] = str(cell.value).strip()
    return headers


def _iter_data_cells(wb):
    """Celdas con contenido, junto a la cabecera de su columna."""
    for ws in wb.worksheets:
        headers = _column_headers(ws)
        header_row = _find_header_row(ws)
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None or _is_formula(cell):
                    continue
                yield ws, cell, headers.get(cell.column, ""), cell.row == header_row


def iter_texts(wb):
    """Unidades de texto del libro, con la cabecera de columna como contexto."""
    for ws, cell, header, _is_header in _iter_data_cells(wb):
        value = cell.value
        if isinstance(value, str) and value.strip():
            yield (value, header)

    for ws in wb.worksheets:
        if ws.title and ws.title.strip():
            yield (ws.title, "")
        for chart in getattr(ws, "_charts", []):
            title = _chart_title_text(chart)
            if title:
                yield (title, "")


def _chart_title_text(chart):
    """Texto del titulo de un grafico, si se puede leer."""
    try:
        title = getattr(chart, "title", None)
        if title is None:
            return ""
        rich = getattr(getattr(title, "tx", None), "rich", None)
        if rich is None:
            return ""
        parts = []
        for paragraph in getattr(rich, "p", []) or []:
            for run in getattr(paragraph, "r", []) or []:
                if getattr(run, "t", None):
                    parts.append(run.t)
        return "".join(parts)
    except Exception:
        return ""


def _numeric_is_sensitive(header: str, policy: int) -> bool:
    if policy == NUMERIC_ALWAYS:
        return True
    if policy == NUMERIC_BY_HEADER:
        return bool(header) and bool(SENSITIVE_HEADER.search(header))
    return False


def _sheet_is_referenced(wb, title: str) -> bool:
    """True si alguna formula del libro nombra esta hoja."""
    escaped = re.escape(title)
    pattern = re.compile(rf"(?:'{escaped}'|{escaped})\s*!")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if _is_formula(cell) and pattern.search(str(cell.value or "")):
                    return True
    return False


def anonymize_excel(input_path: str, output_path: str,
                    lang: str = "es", level: str = "balanced") -> dict:
    wb = openpyxl.load_workbook(input_path)
    engine = PiiEngine(label_lang=lang, level=level)
    engine.prime(iter_texts(wb))

    policy = NUMERIC_POLICY[level_value(level)]
    numeric_changed = 0
    warnings = []

    for ws, cell, header, is_header in _iter_data_cells(wb):
        value = cell.value

        if isinstance(value, str):
            replacements = engine.replacements_for(value, header)
            if replacements:
                cell.value = apply_to_text(value, replacements)
            continue

        # Los valores numericos no llevan contexto propio: la decision depende
        # del nivel y de lo que diga la cabecera de su columna.
        if isinstance(value, bool) or is_header:
            continue
        if isinstance(value, (int, float)) and _numeric_is_sensitive(header, policy):
            cell.value = engine.registry.marker_for("NUMBER", f"{ws.title}!{cell.coordinate}")
            numeric_changed += 1

    # Nombres de hoja: solo si contienen datos personales y nadie los referencia.
    for ws in wb.worksheets:
        title = ws.title or ""
        if not title.strip():
            continue
        anonymized = engine.anonymize_plain(title)
        if anonymized == title:
            continue
        if _sheet_is_referenced(wb, title):
            warnings.append({"code": "sheet_not_renamed", "sheet": title})
            continue
        safe = anonymized.replace("[", "(").replace("]", ")")[:_MAX_SHEET_TITLE]
        try:
            ws.title = safe
        except Exception:
            warnings.append({"code": "sheet_not_renamed", "sheet": title})

    wb.save(output_path)
    stats = engine.stats()
    stats["total"] += numeric_changed
    stats["numeric_cells"] = numeric_changed
    stats["warnings"].extend(warnings)
    return stats
