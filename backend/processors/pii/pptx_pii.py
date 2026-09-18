"""Anonimizacion de presentaciones .pptx conservando el texto no sensible.

Regla dura de este modulo: **no se borra ningun parrafo**. El modo de estructura
elimina todos los parrafos de un cuadro de texto salvo el primero, lo cual es
correcto cuando el contenido se sustituye por una sola etiqueta; aqui destruiria
las vinetas, los saltos de linea y la maqueta del texto que debe conservarse.

La unidad de analisis es el parrafo del ``text_frame``, no el marco entero, para
que los offsets encajen con los runs de ese parrafo.
"""

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .engine import PiiEngine
from .runs import apply_replacements

_TEXT_PROPERTIES = ("author", "last_modified_by", "title", "subject",
                    "comments", "keywords", "category")


def _iter_text_frames(prs):
    """Todos los marcos de texto de la presentacion, descendiendo en grupos."""

    def walk(shapes):
        for shape in shapes:
            try:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from walk(shape.shapes)
                    continue
            except Exception:
                pass                      # forma sin tipo utilizable: se trata como hoja
            try:
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        for cell in row.cells:
                            yield cell.text_frame
                    continue
                if getattr(shape, "has_text_frame", False):
                    yield shape.text_frame
            except Exception:
                continue

    for slide in prs.slides:
        yield from walk(slide.shapes)
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes is not None:
                    yield notes
        except Exception:
            continue


def _iter_paragraphs(prs):
    for frame in _iter_text_frames(prs):
        try:
            for paragraph in frame.paragraphs:
                yield paragraph
        except Exception:
            continue


def iter_texts(prs):
    for paragraph in _iter_paragraphs(prs):
        try:
            runs = paragraph.runs
        except Exception:
            continue
        if not runs:
            continue
        text = "".join(run.text or "" for run in runs)
        if text.strip():
            yield (text, "")


def _anonymize_core_properties(prs, engine) -> int:
    changed = 0
    try:
        props = prs.core_properties
    except Exception:
        return 0
    for name in _TEXT_PROPERTIES:
        try:
            value = getattr(props, name, None)
        except Exception:
            continue
        if not value or not str(value).strip():
            continue
        text = str(value)
        new_value = (engine.registry.marker_for("PERSON", text)
                     if name in ("author", "last_modified_by")
                     else engine.anonymize_plain(text))
        if new_value != text:
            try:
                setattr(props, name, new_value)
                changed += 1
            except Exception:
                continue
    return changed


def anonymize_pptx(input_path: str, output_path: str,
                   lang: str = "es", level: str = "balanced") -> dict:
    prs = Presentation(input_path)
    engine = PiiEngine(label_lang=lang, level=level)
    engine.prime(iter_texts(prs))

    skipped = 0
    for paragraph in _iter_paragraphs(prs):
        try:
            runs = paragraph.runs
            if not runs:
                continue
            segments = [run.text or "" for run in runs]
            text = "".join(segments)
            replacements = engine.replacements_for(text)
            if not replacements:
                continue
            for run, new_text in zip(runs, apply_replacements(segments, replacements)):
                run.text = new_text
        except Exception:
            # En modo PII un parrafo saltado es una fuga, asi que se cuenta y se
            # informa en vez de ignorarlo en silencio.
            skipped += 1
            continue

    meta_changed = _anonymize_core_properties(prs, engine)

    prs.save(output_path)
    stats = engine.stats()
    stats["total"] += meta_changed
    stats["metadata_cleaned"] = meta_changed
    if skipped:
        stats["warnings"].append({"code": "shapes_skipped", "count": str(skipped)})
    return stats
