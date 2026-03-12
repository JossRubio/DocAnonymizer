import math
from lxml import etree
from docx import Document
from docx.oxml.ns import qn

_W_T           = qn('w:t')
_W_R           = qn('w:r')
_W_BR          = qn('w:br')
_W_P           = qn('w:p')
_W_PPR         = qn('w:pPr')
_W_PSTYLE      = qn('w:pStyle')
_W_RPR         = qn('w:rPr')
_W_SZ          = qn('w:sz')
_W_NUMPR       = qn('w:numPr')
_W_TXBXCONTENT = qn('w:txbxContent')
_XML_SPACE     = '{http://www.w3.org/XML/1998/namespace}space'

_BASE_CHARS_PER_LINE_AT_12PT = 80


def _direct_t_elements(p_elem):
    result = []
    for child in p_elem:
        if child.tag == _W_R:
            for sub in child:
                if sub.tag == _W_T:
                    result.append(sub)
    return result


def _direct_text(p_elem) -> str:
    return ''.join(t.text or '' for t in _direct_t_elements(p_elem))


def _get_style_name(p_elem) -> str:
    ppr = p_elem.find(_W_PPR)
    if ppr is not None:
        ps = ppr.find(_W_PSTYLE)
        if ps is not None:
            return (ps.get(qn('w:val')) or '').lower()
    return ''


def _has_numpr(p_elem) -> bool:
    ppr = p_elem.find(_W_PPR)
    return ppr is not None and ppr.find(_W_NUMPR) is not None


def _get_font_halfpt(p_elem) -> int:
    for child in p_elem:
        if child.tag == _W_R:
            rpr = child.find(_W_RPR)
            if rpr is not None:
                sz = rpr.find(_W_SZ)
                if sz is not None:
                    val = sz.get(qn('w:val'))
                    if val:
                        return int(val)
    return 24


def _estimate_lines(text: str, font_halfpt: int) -> int:
    if not text:
        return 1
    chars_per_line = _BASE_CHARS_PER_LINE_AT_12PT * (12.0 / (font_halfpt / 2.0))
    return max(1, math.ceil(len(text) / chars_per_line))


def _replace_paragraph_text(p_elem, new_text: str, original_text: str = ''):
    t_elems = _direct_t_elements(p_elem)
    if not t_elems:
        return

    font_halfpt = _get_font_halfpt(p_elem)
    extra_breaks = max(0,
        _estimate_lines(original_text or new_text, font_halfpt) -
        _estimate_lines(new_text, font_halfpt)
    )

    t_elems[0].text = new_text
    t_elems[0].set(_XML_SPACE, 'preserve')
    for t in t_elems[1:]:
        t.text = ''

    parent_r = t_elems[0].getparent()
    if parent_r is not None and extra_breaks > 0:
        for _ in range(extra_breaks):
            br = etree.SubElement(parent_r, _W_BR)


def process_word(input_path: str, output_path: str) -> dict:
    doc = Document(input_path)
    labels_used = []
    counters = {
        'title': 0, 'section': 0, 'subtitle': 0,
        'paragraph': 0, 'list_item': 0,
    }

    def label(tag: str) -> str:
        labels_used.append(tag)
        return tag

    def classify(p_elem):
        text = _direct_text(p_elem).strip()
        if not text:
            return None
        sn = _get_style_name(p_elem).replace(' ', '').replace('-', '')

        if 'heading1' in sn or sn == 'title':
            if counters['title'] == 0:
                counters['title'] += 1
                return label('[TÍTULO DEL DOCUMENTO]')
            counters['section'] += 1
            return label(f"[TÍTULO SECCIÓN {counters['section']}]")

        if 'heading2' in sn:
            counters['section'] += 1
            return label(f"[TÍTULO SECCIÓN {counters['section']}]")

        if 'heading3' in sn or 'subtitle' in sn:
            counters['subtitle'] += 1
            s = counters['section'] or 1
            return label(f"[SUBTÍTULO {s}.{counters['subtitle']}]")

        if 'heading' in sn:
            counters['section'] += 1
            return label(f"[TÍTULO SECCIÓN {counters['section']}]")

        if 'list' in sn or _has_numpr(p_elem):
            counters['list_item'] += 1
            return label(f"[ELEMENTO LISTA {counters['list_item']}]")

        counters['paragraph'] += 1
        return label(f"[CONTENIDO PÁRRAFO {counters['paragraph']}]")

    for para in doc.paragraphs:
        p = para._element
        orig = _direct_text(p)
        lbl = classify(p)
        if lbl:
            _replace_paragraph_text(p, lbl, orig)

    for table in doc.tables:
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for para in cell.paragraphs:
                    orig = _direct_text(para._element)
                    if orig.strip():
                        lbl = label(f'[CELDA TABLA FILA-{r_idx+1} COL-{c_idx+1}]')
                        _replace_paragraph_text(para._element, lbl, orig)

    txbx_counter = 0
    for txbx in doc.element.body.iter(_W_TXBXCONTENT):
        for p_elem in txbx.iter(_W_P):
            orig = _direct_text(p_elem)
            if orig.strip():
                txbx_counter += 1
                lbl = label(f'[CUADRO DE TEXTO {txbx_counter}]')
                _replace_paragraph_text(p_elem, lbl, orig)

    for section in doc.sections:
        for para in section.header.paragraphs:
            orig = _direct_text(para._element)
            if orig.strip():
                _replace_paragraph_text(para._element, label('[ENCABEZADO]'), orig)
        for para in section.footer.paragraphs:
            orig = _direct_text(para._element)
            if orig.strip():
                _replace_paragraph_text(para._element, label('[PIE DE PÁGINA]'), orig)

    doc.save(output_path)
    return {'total': len(labels_used), 'labels': list(dict.fromkeys(labels_used))}
