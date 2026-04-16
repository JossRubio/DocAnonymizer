from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from processors.labels import get_labels


def _replace_text_in_frame(tf, new_text: str):
    """Replace all text in a text frame, preserving the first run's formatting."""
    paragraphs = tf.paragraphs
    if not paragraphs:
        return
    first_para = paragraphs[0]
    for para in paragraphs[1:]:
        para._p.getparent().remove(para._p)
    runs = first_para.runs
    if runs:
        runs[0].text = new_text
        for run in runs[1:]:
            run.text = ""
    else:
        first_para.add_run().text = new_text


def _process_shape(shape, slide_num, counters, labels_used, L):
    """Process a single shape, recursing into groups."""

    def label(tag):
        labels_used.append(tag)
        return tag

    # ---- GROUP: recurse into sub-shapes ----
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            _process_shape(child, slide_num, counters, labels_used, L)
        return

    # ---- TABLE ----
    if getattr(shape, 'has_table', False):
        for r_idx, row in enumerate(shape.table.rows):
            for c_idx, cell in enumerate(row.cells):
                if cell.text.strip():
                    lbl = label(L['table_cell_pptx'].format(s=slide_num, r=r_idx+1, c=c_idx+1))
                    _replace_text_in_frame(cell.text_frame, lbl)
        return

    # ---- TEXT FRAME ----
    if not getattr(shape, 'has_text_frame', False):
        return

    tf = shape.text_frame
    text = tf.text.strip()
    if not text:
        return

    # Determine label based on placeholder type or shape role
    is_placeholder = getattr(shape, 'is_placeholder', False)
    ph = shape.placeholder_format if is_placeholder else None

    if ph is not None:
        ph_idx = ph.idx
        if ph_idx == 0:
            counters["slide_title"] += 1
            lbl = label(L['slide_title'].format(n=slide_num))
        elif ph_idx == 1:
            counters["body"] += 1
            lbl = label(L['slide_subtitle'].format(n=slide_num))
        else:
            counters["body"] += 1
            lbl = label(L['body_text'].format(n=counters['body']))
    else:
        name_lower = shape.name.lower()
        if "footer" in name_lower or "pie" in name_lower:
            lbl = label(L['slide_footer'])
        else:
            counters["shape"] += 1
            lbl = label(L['shape_label'].format(n=counters['shape']))

    _replace_text_in_frame(tf, lbl)


def _iter_shapes(shapes):
    """Yield all shapes recursively, descending into groups."""
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def process_pptx(input_path: str, output_path: str, lang: str = 'es') -> dict:
    prs = Presentation(input_path)
    L = get_labels(lang)
    labels_used = []
    counters = {
        "slide_title": 0,
        "body": 0,
        "shape": 0,
        "note": 0,
    }

    for slide_num, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            try:
                _process_shape(shape, slide_num, counters, labels_used, L)
            except Exception:
                # Skip malformed shapes without crashing the whole presentation
                pass

        # Presenter notes
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            if notes_tf and notes_tf.text.strip():
                counters["note"] += 1
                lbl = L['presenter_note'].format(n=slide_num)
                labels_used.append(lbl)
                _replace_text_in_frame(notes_tf, lbl)

    prs.save(output_path)
    unique_labels = list(dict.fromkeys(labels_used))
    return {"total": len(labels_used), "labels": unique_labels}
