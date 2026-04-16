import openpyxl
from openpyxl.utils import get_column_letter
from processors.labels import get_labels


def process_excel(input_path: str, output_path: str, lang: str = 'es') -> dict:
    wb = openpyxl.load_workbook(input_path)
    L = get_labels(lang)
    labels_used = []
    counters = {
        "sheet": 0,
        "header": 0,
        "dato": 0,
        "valor": 0,
        "formula": 0,
        "chart": 0,
    }

    def label(tag: str) -> str:
        labels_used.append(tag)
        return tag

    for sheet_idx, ws in enumerate(wb.worksheets, start=1):
        # Rename sheet
        counters["sheet"] += 1
        ws.title = label(L['sheet_name'].format(n=counters['sheet']))

        max_row = ws.max_row or 0
        max_col = ws.max_column or 0

        for row_idx in range(1, max_row + 1):
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                value = cell.value

                if value is None:
                    continue

                str_val = str(value)

                if str_val.startswith("="):
                    counters["formula"] += 1
                    cell.value = label(L['formula'].format(n=counters['formula']))
                elif row_idx == 1:
                    counters["header"] += 1
                    cell.value = label(L['header_col'].format(n=col_idx))
                elif isinstance(value, (int, float)):
                    counters["valor"] += 1
                    cell.value = label(L['numeric_value'].format(n=counters['valor']))
                else:
                    counters["dato"] += 1
                    cell.value = label(L['data_cell'].format(r=row_idx, c=col_idx))

        # Chart titles
        for chart in getattr(ws, '_charts', []):
            counters["chart"] += 1
            lbl = label(L['chart_title'].format(n=counters['chart']))
            if hasattr(chart, "title") and chart.title is not None:
                try:
                    chart.title = lbl
                except Exception:
                    pass

    wb.save(output_path)
    unique_labels = list(dict.fromkeys(labels_used))
    return {"total": len(labels_used), "labels": unique_labels}
